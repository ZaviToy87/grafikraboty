# -*- coding: utf-8 -*-
"""
hr_docs.py — генерация кадровых документов ООО «КАКИЕ ЛЮДИ» (ВетГид) в Word.

Комплект нового сотрудника (2026):
  * Анкета сотрудника
  * Договор о полной индивидуальной материальной ответственности
  * Агентский договор
  * Договор о неразглашении конфиденциальной информации
  * Памятка с учётными данными для входа в программу
Образцы журналов для магазина/ветаптеки:
  * Журнал учёта температуры и влажности (холодильники)
  * График уборки помещений
  * Журнал списания/брака
"""
import os
import io
import json
import re
import zipfile
from datetime import datetime

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FONT = 'Times New Roman'
DOCS_DIR = os.path.join(BASE_DIR, 'docs')
RUS_MONTHS = ['января', 'февраля', 'марта', 'апреля', 'мая', 'июня',
              'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря']


def rus_date_text(d):
    return '«%02d» %s %d г.' % (d.day, RUS_MONTHS[d.month - 1], d.year)


def ensure_dates(ctx):
    """Ставит дату составления = сегодня, если она не заполнена/шаблонная,
    и срок окончания = один год от даты составления."""
    import re as _re
    from datetime import timedelta
    dc = str(ctx.get('date_contract') or '')
    if not dc or '___' in dc or '…' in dc:
        dc = rus_date_text(datetime.now())
        ctx['date_contract'] = dc
    m = _re.search(r'«(\d{1,2})»\s+([а-яА-ЯёЁ]+)\s+(\d{4})', dc)
    if m:
        try:
            day = int(m.group(1))
            year = int(m.group(3))
            mon_name = m.group(2).lower()
            month = next((i for i, name in enumerate(RUS_MONTHS, 1)
                          if name.startswith(mon_name[:6])
                          or mon_name.startswith(name[:6])), 1)
            start = datetime(year, month, day)
            end = start.replace(year=year + 1) - timedelta(days=1)
            ctx['contract_start_text'] = rus_date_text(start)
            ctx['contract_end_text'] = rus_date_text(end)
        except Exception:
            pass
    if not ctx.get('contract_start_text'):
        ctx['contract_start_text'] = dc
        ctx['contract_end_text'] = dc
    return ctx


def load_requisites():
    path = os.path.join(BASE_DIR, 'company_requisites.json')
    try:
        with io.open(path, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return {}


def _split_fio(ctx):
    """Разбивает полное ФИО на фамилию/имя/отчество, если частей нет."""
    if ctx.get('surname') or ctx.get('name') or ctx.get('patronymic'):
        return ctx
    words = [w for w in (ctx.get('fio') or '').strip().split() if w]
    if not words:
        return ctx
    if len(words) >= 3:
        ctx['surname'], ctx['name'], ctx['patronymic'] = words[0], words[1], words[2]
    elif len(words) == 2:
        ctx['surname'], ctx['name'] = words[0], words[1]
    else:
        ctx['name'] = words[0]
    return ctx


def _base_ctx(emp):
    req = load_requisites()
    org = req.get('organization', {})
    ctx = dict(org)
    ctx['year'] = req.get('docs_year', 2026)
    ctx.update(emp or {})
    ctx = _split_fio(ctx)
    return ensure_dates(ctx)


def new_doc(landscape=False):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(1.5)
    if landscape:
        sec.orientation = 1  # landscape
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    return doc


def P(doc, text='', bold=False, center=False, right=False, size=None,
      space_after=4, underline=False, italic=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif right:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    parts = str(text).split('\n')
    for i, part in enumerate(parts):
        r = p.add_run(part)
        r.bold = bold
        r.italic = italic
        r.underline = underline
        if size:
            r.font.size = Pt(size)
        if i < len(parts) - 1:
            r.add_break()
    return p


def line(doc, label, width=50):
    """Строка с подчёркиванием для заполнения."""
    P(doc, '%s %s' % (label, '_' * width), space_after=6)


def value_or_blank(v, width=42):
    v = (v or '').strip()
    return v if v else '_' * width


def fill(text, ctx):
    import re as _re
    def rep(m):
        return str(ctx.get(m.group(1), ''))
    return _re.sub(r'\{([A-Za-z0-9_]+)\}', rep, text)


def org_header(doc, ctx, city='Самара'):
    """Шапка документа: реквизиты берём из конфига ООО (не из данных сотрудника)."""
    org = load_requisites().get('organization', {})
    P(doc, org.get('full_name', ''), center=True, bold=True, size=13)
    P(doc, 'Юр. адрес: %s' % org.get('legal_address', ''), center=True, size=10)
    P(doc, 'Факт. адрес: %s' % org.get('actual_address', ''), center=True,
      size=10)
    phone = org.get('phone', '')
    email = org.get('email', '')
    if phone or email:
        P(doc, ('Тел.: %s    E-mail: %s' % (phone, email)).strip(),
          center=True, size=10)
    P(doc, 'ИНН %s, КПП %s, ОГРН %s' % (org.get('inn', ''),
                                         org.get('kpp', ''),
                                         org.get('ogrn', '')),
      center=True, size=10)
    P(doc, '', space_after=4)


def place_and_date(doc, city='г. Тольятти', date_text=''):
    P(doc, city, size=12, space_after=0)
    P(doc, date_text or '«___» ____________ 2026 г.', right=True, size=12,
      space_after=8)


def add_heading(doc, text, size=13):
    P(doc, text, center=True, bold=True, size=size, space_after=2)


def _cell_lines(cell, lines, bold_first=False):
    cell.text = ''
    first = True
    for ln in lines:
        if ln == '':
            if not first:
                cell.paragraphs[-1].add_run('\n')
            continue
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.1
        r = p.add_run(str(ln))
        r.bold = bold_first and first
        r.font.size = Pt(11)
        first = False


def render_parties(doc, ctx, worker_label='Работник'):
    """Блок «Реквизиты и подписи сторон» в виде аккуратной таблицы."""
    org = load_requisites().get('organization', {})
    P(doc, '', space_after=2)
    P(doc, 'РЕКВИЗИТЫ И ПОДПИСИ СТОРОН', center=True, bold=True, size=12,
      space_after=6)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = 'Table Grid'
    left = [
        'РАБОТОДАТЕЛЬ',
        org.get('full_name', ''),
        'ОГРН %s' % org.get('ogrn', ''),
        'ИНН %s, КПП %s' % (org.get('inn', ''), org.get('kpp', '')),
        'Юр. адрес: %s' % org.get('legal_address', ''),
        'Факт. адрес: %s' % org.get('actual_address', ''),
        'Банк: %s' % org.get('bank_name', ''),
        'БИК %s, р/с %s' % (org.get('bank_bik', ''),
                            org.get('bank_account', '')),
        'к/с %s' % org.get('bank_corr_account', ''),
        '',
        fill('{director_position}:', ctx),
        '______________________ /{director_fio}/'.format(**ctx),
        'М.П.',
    ]
    right = [
        worker_label.upper(),
        ctx.get('fio') or '',
        'Должность: %s' % (ctx.get('position') or '____________'),
    ]
    if (ctx.get('birth_date') or '').strip():
        right.append('Дата рождения: %s' % ctx.get('birth_date'))
    if (ctx.get('phone') or '').strip():
        right.append('Телефон: %s' % ctx.get('phone'))
    if (ctx.get('inn') or '').strip():
        right.append('ИНН: %s' % ctx.get('inn'))
    if (ctx.get('snils') or '').strip():
        right.append('СНИЛС: %s' % ctx.get('snils'))
    right += [
        '',
        'Паспорт: %s' % (ctx.get('passport') or '________________________'),
        'Кем выдан: %s' % (ctx.get('passport_by') or '____________________'),
        'Дата выдачи: %s' % (ctx.get('passport_date') or '______________'),
        'Код подразделения: %s' % (ctx.get('passport_code') or '________'),
        '',
        'Адрес регистрации:',
        ctx.get('address_registration') or '____________________________',
        '',
        'Подпись: ______________________',
    ]
    _cell_lines(tbl.cell(0, 0), left, bold_first=True)
    _cell_lines(tbl.cell(0, 1), right, bold_first=True)
    return tbl


def build_anketa(emp, path):
    """Анкета сотрудника."""
    ctx = _base_ctx(emp)
    doc = new_doc()
    P(doc, 'АНКЕТА СОТРУДНИКА', center=True, bold=True, size=14, space_after=2)
    P(doc, ctx.get('short_name', '') + ' (бренд %s)' % ctx.get('brand', ''),
      center=True, size=11, space_after=8)
    P(doc, 'Заполняется при приёме на работу.', size=10, italic=True)
    fields = [
        ('Фамилия', 'surname'), ('Имя', 'name'), ('Отчество', 'patronymic'),
        ('Дата рождения', 'birth_date'), ('Телефон', 'phone'),
        ('Адрес регистрации', 'address_registration'),
        ('Адрес проживания', 'address_residence'),
        ('Серия и номер паспорта', 'passport'),
        ('Кем выдан паспорт', 'passport_by'),
        ('Дата выдачи', 'passport_date'), ('Код подразделения', 'passport_code'),
        ('ИНН', 'inn'), ('СНИЛС', 'snils'),
        ('Должность', 'position'), ('Дата начала работы', 'date_start'),
        ('График работы', 'work_schedule'),
        ('Заработная плата (оклад/порядок)', 'salary_scheme'),
    ]
    for label, key in fields:
        val = (ctx.get(key) or '').strip()
        if val:
            P(doc, '%s: %s' % (label, val), space_after=6)
        else:
            line(doc, label + ':', 60)
    P(doc, '', space_after=4)
    P(doc, 'С правилами внутреннего трудового распорядка, должностной '
           'инструкцией, правилами охраны труда и пожарной безопасности, '
           'Методичкой сотрудника ознакомлен(а):', size=11)
    P(doc, 'Подпись ____________________ / %s /  %s г.'
      % (value_or_blank(ctx.get('fio'), 34),
         value_or_blank(ctx.get('date_start', ctx.get('year', '')), 18)), size=11)
    doc.save(path)
    return path


def build_material_contract(emp, path):
    """Договор о полной индивидуальной материальной ответственности."""
    ctx = _base_ctx(emp)
    doc = new_doc()
    org_header(doc, ctx)
    ctx['doc_number'] = next_doc_number('material')
    P(doc, 'ДОГОВОР № %s' % ctx['doc_number'], center=True, bold=True,
      size=13, space_after=0)
    P(doc, 'о полной индивидуальной материальной ответственности работника',
      center=True, bold=True, size=13, space_after=2)
    place_and_date(doc, 'г. Тольятти', ctx.get('date_contract', ''))
    intro = fill('{short_name} (далее — «Работодатель»), в лице {director_of} '
                 '{director_fio_gen}, действующего на основании {director_basis}, '
                 'с одной стороны, и гражданин(ка) РФ {fio} (далее — «Работник»), '
                 'с другой стороны, заключили настоящий договор о нижеследующем:', ctx)
    P(doc, intro)
    P(doc, '1. ПРЕДМЕТ ДОГОВОРА', bold=True)
    P(doc, '1.1. Работник, занимающий должность «{position}», принимает на '
           'себя полную индивидуальную материальную ответственность за '
           'недостачу вверенного ему Работодателем имущества, а также за '
           'ущерб, возникший у Работодателя в результате возмещения им '
           'ущерба иным лицам.'.format(**ctx))
    P(doc, '1.2. Настоящий договор заключён в соответствии со ст. 242–244 '
           'Трудового кодекса РФ и Перечнем должностей и работ, замещаемых '
           'или выполняемых работниками, с которыми работодатель может '
           'заключать письменные договоры о полной индивидуальной '
           'материальной ответственности, утверждённым Постановлением '
           'Минтруда России от 31.12.2002 № 85.')
    P(doc, '2. ОБЯЗАННОСТИ РАБОТНИКА', bold=True)
    P(doc, '2.1. Бережно относиться к переданному ему имуществу '
           'Работодателя и принимать меры к предотвращению ущерба.')
    P(doc, '2.2. Своевременно сообщать Работодателю обо всех '
           'обстоятельствах, угрожающих обеспечению сохранности вверенного '
           'имущества.')
    P(doc, '2.3. Вести учёт, своевременно и в установленном порядке '
           'составлять и представлять отчёты о движении и остатках '
           'вверенного имущества (товаров, денежных средств).')
    P(doc, '2.4. Участвовать в проведении инвентаризаций, ревизий и иных '
           'проверок сохранности имущества.')
    P(doc, '3. ОБЯЗАННОСТИ РАБОТОДАТЕЛЯ', bold=True)
    P(doc, '3.1. Создавать Работнику условия, необходимые для нормальной '
           'работы и обеспечения полной сохранности вверенного имущества.')
    P(doc, '3.2. Своевременно проводить инвентаризации и ревизии, '
           'фиксировать их результаты.')
    P(doc, '3.3. Знакомить Работника с действующим законодательством о '
           'материальной ответственности.')
    P(doc, '4. ПОРЯДОК ОПРЕДЕЛЕНИЯ РАЗМЕРА УЩЕРБА', bold=True)
    P(doc, '4.1. Размер ущерба определяется по фактическим потерям на '
           'основании данных бухгалтерского учёта (ст. 246 ТК РФ). При '
           'недостаче имущества размер ущерба определяется из рыночной '
           'стоимости на день причинения ущерба, но не ниже балансовой '
           'стоимости по данным бухгалтерского учёта.')
    P(doc, '5. ОТВЕТСТВЕННОСТЬ СТОРОН', bold=True)
    P(doc, '5.1. Работник несёт полную материальную ответственность за '
           'недостачу вверенного имущества в размере причинённого ущерба '
           '(ст. 242, 243 ТК РФ).')
    P(doc, '5.2. Работник не несёт материальной ответственности, если '
           'ущерб возник вследствие непреодолимой силы, нормального '
           'хозяйственного риска, крайней необходимости или необходимой '
           'обороны либо неисполнения Работодателем обязанности по '
           'обеспечению надлежащих условий для хранения вверенного '
           'имущества (ст. 239 ТК РФ).')
    P(doc, '6. ПРОЧИЕ УСЛОВИЯ', bold=True)
    P(doc, '6.1. Договор вступает в силу с момента подписания и действует '
           'в течение всего периода работы Работника с вверенным ему '
           'имуществом.')
    P(doc, '6.2. Договор составлен в двух экземплярах, имеющих одинаковую '
           'юридическую силу: один — у Работодателя, второй — у Работника.')
    render_parties(doc, ctx)
    doc.save(path)
    return path


def build_agent_contract(emp, path):
    """Агентский договор с сотрудником (продавцом-агентом)."""
    ctx = _base_ctx(emp)
    doc = new_doc()
    org_header(doc, ctx)
    P(doc, 'АГЕНТСКИЙ ДОГОВОР № ______', center=True, bold=True, size=13,
      space_after=0)
    P(doc, 'на совершение действий от имени Принципала', center=True,
      bold=True, size=12, space_after=2)
    place_and_date(doc, 'г. Тольятти', ctx.get('date_contract', ''))
    intro = fill('{short_name} (далее — «Принципал»), в лице {director_of} '
                 '{director_fio_gen}, действующего на основании '
                 '{director_basis}, с одной стороны, и {fio} (далее — «Агент»), '
                 'с другой стороны, заключили настоящий договор о '
                 'нижеследующем:', ctx)
    P(doc, intro)
    P(doc, '1. ПРЕДМЕТ ДОГОВОРА', bold=True)
    P(doc, '1.1. По настоящему договору Агент обязуется от имени и за счёт '
           'Принципала совершать юридические и иные действия по продаже '
           'товаров зоомагазина и ветаптеки (бренд «{brand}»), приёму '
           'денежных средств от покупателей, ведению кассовых операций и '
           'предоставлению покупателям информации о товарах.'.format(**ctx))
    P(doc, '1.2. Агент действует на территории г. Тольятти и '
           'Самарской области.')
    P(doc, '1.3. Полномочия Агента подтверждаются доверенностью, '
           'выдаваемой Принципалом.')
    P(doc, '2. ПРАВА И ОБЯЗАННОСТИ АГЕНТА', bold=True)
    P(doc, '2.1. Агент обязуется:')
    P(doc, '— лично совершать действия, предусмотренные п. 1.1 договора;')
    P(doc, '— соблюдать правила внутреннего распорядка, Методичку '
           'сотрудника и инструкции Принципала;')
    P(doc, '— обеспечивать сохранность товара и денежных средств, '
           'своевременно сдавать выручку;')
    P(doc, '— не разглашать конфиденциальную информацию Принципала.')
    P(doc, '2.2. Агент вправе получать вознаграждение в порядке, '
           'предусмотренном разделом 3 договора.')
    P(doc, '3. ВОЗНАГРАЖДЕНИЕ АГЕНТА', bold=True)
    P(doc, '3.1. Вознаграждение Агента устанавливается сторонами в '
           'размере, согласованном в дополнительном соглашении / листе '
           'согласования (процент от выручки или фиксированная оплата за '
           'смену по действующей у Принципала схеме мотивации).')
    P(doc, '3.2. Вознаграждение выплачивается не реже одного раза в месяц '
           'в порядке, установленном локальными актами Принципала.')
    P(doc, '4. ОТЧЁТНОСТЬ АГЕНТА', bold=True)
    P(doc, '4.1. Агент предоставляет Принципалу отчёт о выполненной работе '
           'в порядке и сроки, установленные Принципалом (по итогам '
           'каждой смены).')
    P(doc, '5. ОТВЕТСТВЕННОСТЬ СТОРОН', bold=True)
    P(doc, '5.1. За неисполнение или ненадлежащее исполнение обязанностей '
           'по договору стороны несут ответственность в соответствии с '
           'законодательством РФ.')
    P(doc, '5.2. Агент несёт материальную ответственность за утрату '
           '(недостачу) вверенных товарно-материальных ценностей в '
           'порядке ст. 242–244 ТК РФ и отдельного договора о полной '
           'материальной ответственности.')
    P(doc, '6. СРОК ДЕЙСТВИЯ И ПРОЧИЕ УСЛОВИЯ', bold=True)
    P(doc, '6.1. Договор заключён в соответствии с главой 52 ГК РФ '
           '(ст. 1005–1011) и действует с даты подписания по 31.12.%s г.'
           % ctx.get('year', 2026))
    P(doc, '6.2. Договор может быть расторгнут по основаниям, '
           'предусмотренным законодательством РФ.')
    P(doc, '6.3. Договор составлен в двух экземплярах, имеющих одинаковую '
           'юридическую силу.')
    render_parties(doc, ctx, 'Агент')
    doc.save(path)
    return path


def build_nda_contract(emp, path):
    """Договор о неразглашении конфиденциальной информации."""
    ctx = _base_ctx(emp)
    doc = new_doc()
    org_header(doc, ctx)
    ctx['doc_number'] = next_doc_number('nda')
    P(doc, 'ДОГОВОР № %s' % ctx['doc_number'], center=True, bold=True,
      size=13, space_after=0)
    P(doc, 'о неразглашении конфиденциальной информации '
           '(коммерческой тайны)', center=True, bold=True, size=12,
      space_after=2)
    place_and_date(doc, 'г. Тольятти', ctx.get('date_contract', ''))
    intro = fill('{short_name} (далее — «Компания»), в лице {director_of} '
                 '{director_fio_gen}, действующего на основании '
                 '{director_basis}, с одной стороны, и {fio} (далее — '
                 '«Сотрудник»), с другой стороны, заключили настоящий '
                 'договор о нижеследующем:', ctx)
    P(doc, intro)
    P(doc, '1. ОПРЕДЕЛЕНИЯ И ПРЕДМЕТ', bold=True)
    P(doc, '1.1. Конфиденциальная информация (КИ) — сведения, составляющие '
           'коммерческую тайну Компании, а также иные сведения, ставшие '
           'известными Сотруднику в связи с выполнением трудовых/'
           'гражданско-правовых обязанностей: о покупателях и '
           'контрагентах, ценах и условиях закупок, выручке, учётных '
           'данных и паролях, внутренних документах, методиках продаж.')
    P(doc, '1.2. Сотрудник обязуется не разглашать КИ третьим лицам и не '
           'использовать её в личных целях в период действия договора и в '
           'течение 3 (трёх) лет после его прекращения.')
    P(doc, '2. ОБЯЗАННОСТИ СТОРОН', bold=True)
    P(doc, '2.1. Компания обязана:')
    P(doc, '— ознакомить Сотрудника с перечнем сведений, составляющих '
           'коммерческую тайну;')
    P(doc, '— обеспечить соблюдение режима коммерческой тайны.')
    P(doc, '2.2. Сотрудник обязан:')
    P(doc, '— не передавать КИ третьим лицам без письменного согласия '
           'Компании;')
    P(doc, '— не использовать КИ в личных целях и не допускать её '
           'разглашения, включая публикации в сети Интернет и мессенджерах;')
    P(doc, '— соблюдать режим доступа к учётным записям программы, не '
           'передавать свой пароль другим лицам;')
    P(doc, '— незамедлительно сообщать Компании о фактах возможного '
           'разглашения КИ.')
    P(doc, '3. ОТВЕТСТВЕННОСТЬ', bold=True)
    P(doc, '3.1. За разглашение КИ Сотрудник несёт ответственность в '
           'соответствии с законодательством РФ, включая возмещение '
           'причинённых убытков.')
    P(doc, '3.2. Отношения по охране коммерческой тайны регулируются '
           'Федеральным законом от 29.07.2004 № 98-ФЗ «О коммерческой '
           'тайне», ст. 1465–1472 ГК РФ, а в части персональных данных — '
           'Федеральным законом от 27.07.2006 № 152-ФЗ «О персональных '
           'данных».')
    P(doc, '4. СРОК ДЕЙСТВИЯ', bold=True)
    P(doc, '4.1. Договор вступает в силу с момента подписания и действует '
           'в течение всего периода работы Сотрудника, а обязательства о '
           'неразглашении КИ — в течение 3 (трёх) лет после прекращения '
           'договора.')
    P(doc, '4.2. Договор составлен в двух экземплярах, имеющих одинаковую '
           'юридическую силу.')
    render_parties(doc, ctx, 'Сотрудник')
    doc.save(path)
    return path


def build_access_memo(emp, path):
    """Памятка сотрудника: учётные данные и первые шаги."""
    ctx = _base_ctx(emp)
    doc = new_doc()
    P(doc, 'ПАМЯТКА СОТРУДНИКА (учётные данные и первые шаги)',
      center=True, bold=True, size=13, space_after=6)
    P(doc, ctx.get('short_name', '') + ' — бренд ' + ctx.get('brand', ''),
      center=True, size=11, space_after=8)
    P(doc, fill('Сотрудник: {fio}', ctx))
    P(doc, fill('Должность: {position}', ctx))
    P(doc, fill('Дата начала работы: {date_start}', ctx))
    P(doc, '', space_after=4)
    P(doc, 'ДОСТУП К ПРОГРАММЕ «ВЕТГИД ГРАФИК»:', bold=True)
    P(doc, fill('Адрес входа: {site_url}', ctx))
    P(doc, fill('Логин: {login}', ctx))
    P(doc, fill('Пароль: {password}', ctx))
    P(doc, '', space_after=4)
    P(doc, 'ВАЖНО:', bold=True)
    P(doc, '• Никому не сообщайте свой пароль. За разглашение пароля и '
           'доступа предусмотрена ответственность по договору о '
           'неразглашении.')
    P(doc, '• Перед первой сменой прочитайте «Методичку сотрудника» и '
           '«Правила магазина» (раздел «📖 Правила и методичка» в '
           'программе).')
    P(doc, '• В начале смены откройте кассовую смену в программе, в конце '
           '— закройте и сдайте выручку по инструкции.')
    P(doc, '• При возникновении вопросов обращайтесь к руководству.')
    P(doc, '', space_after=4)
    P(doc, 'Ознакомлен(а) и обязуюсь соблюдать:')
    P(doc, 'Подпись ____________________ / {fio} / {date_start}'
      .format(**ctx), size=11)
    doc.save(path)
    return path


# ============================================================
# Образцы журналов для магазина и ветаптеки
# ============================================================
def _make_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = t.cell(0, j)
        cell.text = h
        for par in cell.paragraphs:
            for r in par.runs:
                r.bold = True
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            t.cell(i + 1, j).text = val
    return t


def _journal_blank_rows(n, first_col_shift=1):
    return [[''] * first_col_shift + [''] for _ in range(n)]


def build_temperature_journal(path, rows=31):
    doc = new_doc(landscape=True)
    P(doc, 'ЖУРНАЛ', center=True, bold=True, size=14, space_after=0)
    P(doc, 'учёта температуры и влажности в холодильном оборудовании',
      center=True, bold=True, size=12, space_after=6)
    P(doc, 'ООО «КАКИЕ ЛЮДИ» (ВетГид). Холодильное оборудование магазина '
           'и ветаптеки.', size=10)
    P(doc, 'Норма: температура в холодильниках +2…+6 °C, в морозильных '
           'камерах −18 °C и ниже. Замеры проводятся 2 раза в день.',
      size=10, italic=True)
    headers = ['Дата', 'Время', 'Холодильник №1, °C', 'Холодильник №2, °C',
               'Морозильная камера, °C', 'Влажность, %',
               'Отклонение/примечания', 'Подпись']
    rows_data = [[''] * len(headers) for _ in range(rows)]
    _make_table(doc, headers, rows_data)
    doc.save(path)
    return path


def build_cleaning_schedule(path, rows=31):
    doc = new_doc()
    P(doc, 'ГРАФИК УБОРКИ ПОМЕЩЕНИЙ', center=True, bold=True, size=14,
      space_after=2)
    P(doc, 'ООО «КАКИЕ ЛЮДИ» (ВетГид). Месяц: ____________ 2026 г.',
      center=True, size=11, space_after=6)
    P(doc, 'Уборка выполняется ежедневно по «Правилам магазина»; '
           'генеральная уборка — по графику.', size=10, italic=True)
    headers = ['Дата', 'Торговый зал (пол)', 'Полки/витрины',
               'Санузел', 'Служебное помещение', 'Подпись']
    rows_data = [[''] * len(headers) for _ in range(rows)]
    _make_table(doc, headers, rows_data)
    doc.save(path)
    return path


def build_writeoff_journal(path, rows=31):
    doc = new_doc()
    P(doc, 'ЖУРНАЛ СПИСАНИЯ И БРАКА ТОВАРА', center=True, bold=True,
      size=14, space_after=6)
    P(doc, 'ООО «КАКИЕ ЛЮДИ» (ВетГид).', center=True, size=11)
    P(doc, 'Записи вносятся при выявлении брака, повреждения упаковки, '
           'просрочки товара.', size=10, italic=True)
    headers = ['Дата', 'Наименование товара', 'Штрихкод/артикул',
               'Кол-во', 'Причина списания', 'Номер акта', 'Подпись']
    rows_data = [[''] * len(headers) for _ in range(rows)]
    _make_table(doc, headers, rows_data)
    doc.save(path)
    return path


def build_journals(out_dir):
    os.makedirs(out_dir, exist_ok=True)
    files = [
        build_temperature_journal(os.path.join(
            out_dir, 'Журнал_температуры_и_влажности.docx')),
        build_cleaning_schedule(os.path.join(
            out_dir, 'График_уборки_помещений.docx')),
        build_writeoff_journal(os.path.join(
            out_dir, 'Журнал_списания_и_брака.docx')),
    ]
    return files


# ============================================================
# Методичка и правила (входят в пакет каждого сотрудника)
# ============================================================
def build_methodichka(emp, path):
    return _text_doc(path, 'МЕТОДИЧКА СОТРУДНИКА',
                     os.path.join(DOCS_DIR, 'МЕТОДИЧКА_сотрудника.txt'))


def build_pravila(emp, path):
    return _text_doc(path, 'ПРАВИЛА МАГАЗИНА',
                     os.path.join(DOCS_DIR, 'Правила_магазина.txt'))


def _text_doc(path, title, src_path):
    ctx = _base_ctx({})
    doc = new_doc()
    P(doc, ctx.get('short_name', '') + ' — бренд ' + ctx.get('brand', ''),
      center=True, size=10)
    P(doc, title, center=True, bold=True, size=14, space_after=6)
    try:
        with io.open(src_path, 'r', encoding='utf-8') as f:
            lines = f.read().splitlines()
    except Exception:
        lines = []
    import re as _re
    for ln in lines:
        s = ln.strip()
        if not s:
            continue
        bold = bool(_re.match(r'^\s*(\d+\.|[-—•*]?\s*[А-ЯЁA-Z][А-ЯЁA-Z ]{6,})', s)
                     or _re.match(r'^(КАК|НЕ|ДОЛЖЕН|ОБЯЗАН|ПРАВИЛ|МЕТОД)', s,
                                  _re.I))
        P(doc, s, bold=bold, size=11, space_after=3)
    P(doc, '', space_after=2)
    P(doc, 'С документом ознакомлен(а): ______________ / ____________ / 20__ г.',
      size=10)
    doc.save(path)
    return path


# ============================================================
# Пакет документов сотрудника
# ============================================================
# (список DOC_BUILDERS определён ниже — после всех функций-генераторов)

def sanitize(name):
    return re.sub(r'[\\/:*?"<>|]+', '_', str(name or '').strip()) or 'документы'


def default_employee():
    """Пустой (чистовой) профиль сотрудника — заполняется при приёме."""
    return {
        'fio': '', 'surname': '', 'name': '', 'patronymic': '',
        'birth_date': '', 'phone': '', 'passport': '',
        'passport_by': '', 'passport_date': '', 'passport_code': '',
        'address_registration': '', 'address_residence': '',
        'inn': '', 'snils': '', 'position': 'Продавец-кассир',
        'date_start': '', 'date_contract': '«___» ____________ 2026 г.',
        'work_schedule': '', 'salary_scheme': '',
        'login': '', 'password': '', 'site_url': 'http://<адрес сервера>:8080',
    }


DOC_NUM_PREFIX = {'agent': 'А', 'material': 'МО', 'nda': 'НД'}


def next_doc_number(doc_type):
    """Сквозная нумерация договоров по году: А-2026-001, МО-2026-001, НД-2026-001…"""
    import sqlite3
    prefix = DOC_NUM_PREFIX.get(doc_type, 'ДОК')
    year = datetime.now().year
    db_path = os.path.join(BASE_DIR, 'schedule.db')
    try:
        conn = sqlite3.connect(db_path)
        cur = conn.cursor()
        cur.execute('''
            CREATE TABLE IF NOT EXISTS hr_doc_numbers (
                doc_type TEXT, year INTEGER, seq INTEGER,
                PRIMARY KEY (doc_type, year))
        ''')
        row = cur.execute('SELECT seq FROM hr_doc_numbers WHERE doc_type=? '
                          'AND year=?', (doc_type, year)).fetchone()
        seq = (row[0] if row else 0) + 1
        cur.execute('INSERT OR REPLACE INTO hr_doc_numbers (doc_type, year, seq) '
                    'VALUES (?, ?, ?)', (doc_type, year, seq))
        conn.commit()
        conn.close()
        return '%s-%d-%03d' % (prefix, year, seq)
    except Exception:
        return '%s-%d-001' % (prefix, year)


def build_employee_kit(emp, out_dir):
    """Формирует комплект документов нового сотрудника и возвращает пути."""
    ctx = dict(default_employee())
    ctx.update(emp or {})
    os.makedirs(out_dir, exist_ok=True)
    name = sanitize(ctx.get('fio') or 'новый_сотрудник')
    files = []
    for kind, label, fn in DOC_BUILDERS:
        p = os.path.join(out_dir, '%s_%s.docx' % (label, name))
        fn(ctx, p)
        files.append(p)
    return files


def build_kit_zip(emp, zip_path):
    """Формирует пакет и упаковывает его в zip-архив."""
    tmp = zip_path + '.tmp_dir'
    files = build_employee_kit(emp, tmp)
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as z:
        for f in files:
            z.write(f, os.path.basename(f))
    for f in files:
        try:
            os.remove(f)
        except OSError:
            pass
    try:
        os.rmdir(tmp)
    except OSError:
        pass
    return zip_path


def build_agent_contract_full(emp, path):
    """Агентский договор (полная редакция) с продавцом-агентом."""
    ctx = _base_ctx(emp)
    doc = new_doc()
    org_header(doc, ctx)
    ctx['doc_number'] = next_doc_number('agent')
    P(doc, 'АГЕНТСКИЙ ДОГОВОР № %s' % ctx['doc_number'], center=True,
      bold=True, size=13, space_after=0)
    P(doc, 'на совершение юридических и фактических действий по продаже '
           'товаров от имени и за счёт Принципала', center=True, bold=True,
      size=11, space_after=2)
    place_and_date(doc, 'г. Тольятти', ctx.get('date_contract', ''))
    intro = fill('{short_name} (далее — «Принципал»), в лице {director_of} '
                 '{director_fio_gen}, действующего на основании '
                 '{director_basis}, с одной стороны, и {fio} (далее — «Агент»), '
                 'с другой стороны, совместно именуемые «Стороны», заключили '
                 'настоящий агентский договор (далее — «Договор») о '
                 'нижеследующем:', ctx)
    P(doc, intro)

    P(doc, '1. ПРЕДМЕТ ДОГОВОРА', bold=True)
    P(doc, '1.1. Принципал поручает, а Агент принимает на себя обязанность '
           'от имени и за счёт Принципала совершать юридические и '
           'фактические действия по организации розничной продажи товаров '
           'зоомагазина и ветаптеки под брендом «{brand}» (далее — '
           '«Товары»), в том числе:'.format(**ctx))
    P(doc, '— обслуживание покупателей, консультирование по ассортименту и '
           'свойствам Товаров;')
    P(doc, '— приём денежных средств от покупателей и передачу их '
           'Принципалу в установленном порядке;')
    P(doc, '— участие в приёмке, размещении и учёте Товаров;')
    P(doc, '— соблюдение правил торговли, ценовой политики и инструкций '
           'Принципала.')
    P(doc, '1.2. Территория исполнения Договора — г. Тольятти и иные '
           'торговые точки Принципала в Самарской области.')
    P(doc, '1.3. Полномочия Агента подтверждаются доверенностью, '
           'выдаваемой Принципалом.')
    P(doc, '2. ПРАВА И ОБЯЗАННОСТИ АГЕНТА', bold=True)
    P(doc, '2.1. Агент обязуется:')
    P(doc, '— лично и добросовестно исполнять поручения Принципала;')
    P(doc, '— соблюдать требования законодательства о защите прав '
           'потребителей и правила розничной торговли;')
    P(doc, '— обеспечивать сохранность Товаров, денежных средств и '
           'оборудования, своевременно сдавать выручку;')
    P(doc, '— соблюдать Правила магазина, Методичку сотрудника и '
           'трудовую дисциплину;')
    P(doc, '— не разглашать конфиденциальную информацию Принципала;')
    P(doc, '— незамедлительно сообщать Принципалу о нарушениях, '
           'недостачах, повреждениях Товаров и иных обстоятельствах, '
           'влияющих на исполнение Договора.')
    P(doc, '2.2. Агент вправе получать вознаграждение в порядке, '
           'предусмотренном разделом 4 Договора.')
    P(doc, '3. ПРАВА И ОБЯЗАННОСТИ ПРИНЦИПАЛА', bold=True)
    P(doc, '3.1. Принципал обязуется:')
    P(doc, '— передавать Агенту Товары и документы, необходимые для '
           'исполнения Договора;')
    P(doc, '— обеспечивать условия для приёмки и хранения Товаров;')
    P(doc, '— выплачивать Агенту вознаграждение в установленном порядке.')
    P(doc, '4. ВОЗНАГРАЖДЕНИЕ АГЕНТА И ПОРЯДОК РАСЧЁТОВ', bold=True)
    P(doc, '4.1. За исполнение поручения Агент получает вознаграждение '
           'согласно действующей у Принципала системе оплаты труда '
           '(оплата за смену по порогам выручки / процент от продаж), '
           'доведённой до Агента под роспись.')
    P(doc, '4.2. Вознаграждение выплачивается не реже одного раза в месяц '
           'в сроки выплаты заработной платы, установленные у Принципала.')
    P(doc, '5. ОТЧЁТНОСТЬ И КОНТРОЛЬ', bold=True)
    P(doc, '5.1. Агент предоставляет Принципалу отчёты о выполненной '
           'работе (по итогам каждой смены) в порядке, установленном '
           'Принципалом, вместе с документами, подтверждающими передачу '
           'выручки.')
    P(doc, '6. МАТЕРИАЛЬНАЯ ОТВЕТСТВЕННОСТЬ', bold=True)
    P(doc, '6.1. Агент несёт полную индивидуальную материальную '
           'ответственность за недостачу вверенного имущества в '
           'соответствии со ст. 242–244 ТК РФ и отдельным договором о '
           'полной материальной ответственности.')
    P(doc, '7. КОНФИДЕНЦИАЛЬНОСТЬ', bold=True)
    P(doc, '7.1. Стороны обязуются не разглашать конфиденциальную '
           'информацию, ставшую им известной при исполнении Договора, в '
           'том числе сведения о покупателях, ценах, выручке и учётных '
           'данных (ФЗ от 29.07.2004 № 98-ФЗ «О коммерческой тайне»).')
    P(doc, '8. СРОК ДЕЙСТВИЯ И РАСТОРЖЕНИЕ', bold=True)
    P(doc, '8.1. Договор вступает в силу с даты подписания (не ранее '
           'указанной в нём даты составления {contract_start_text}) и '
           'действует по {contract_end_text} включительно (один год) с '
           'возможностью пролонгации по соглашению сторон.'.format(**ctx))
    P(doc, '8.2. Каждая из Сторон вправе отказаться от Договора, '
           'предупредив другую Сторону не позднее чем за 14 (четырнадцать) '
           'календарных дней (ст. 1011 ГК РФ во взаимосвязи со ст. 1005 ГК РФ).')
    P(doc, '9. ЗАКЛЮЧИТЕЛЬНЫЕ ПОЛОЖЕНИЯ', bold=True)
    P(doc, '9.1. Договор заключён в соответствии с главой 52 ГК РФ '
           '(ст. 1005–1011 ГК РФ).')
    P(doc, '9.2. Споры по Договору разрешаются путём переговоров, при '
           'недостижении согласия — в судебном порядке по месту нахождения '
           'Принципала.')
    P(doc, '9.3. Изменения и дополнения к Договору действительны при '
           'условии их оформления в письменной форме и подписания обеими '
           'Сторонами.')
    P(doc, '9.4. Договор составлен в двух экземплярах, имеющих одинаковую '
           'юридическую силу: по одному для каждой из Сторон.')
    render_parties(doc, ctx, 'Агент')
    doc.save(path)
    return path


# Список документов пакета (после всех генераторов)
DOC_BUILDERS = [
    ('anketa', '01_Анкета_сотрудника', build_anketa),
    ('material', '02_Договор_материальная_ответственность',
     build_material_contract),
    ('agent', '03_Агентский_договор', build_agent_contract_full),
    ('nda', '04_Договор_о_неразглашении', build_nda_contract),
    ('memo', '05_Памятка_учётные_данные', build_access_memo),
    ('methodichka', '06_Методичка_сотрудника', build_methodichka),
    ('pravila', '07_Правила_магазина', build_pravila),
]
DOC_LABELS = {k: l for k, l, _ in DOC_BUILDERS}






