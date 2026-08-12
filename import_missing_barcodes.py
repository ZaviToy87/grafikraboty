# -*- coding: utf-8 -*-
"""
Скрипт для импорта недостающих штрих-кодов из Excel
"""
import openpyxl
import sqlite3
import sys

EXCEL_PATH = r'C:\Users\User\Desktop\Новые штрих кода.1.xlsx'
DB_PATH = 'schedule.db'

def main():
    # Загружаем Excel
    print(f"📖 Загружаем Excel: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    
    # Считываем товары из Excel
    excel_products = {}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        name = row[0].value if row[0].value else ''
        factory = row[1].value if len(row) > 1 and row[1].value else ''
        internal = row[2].value if len(row) > 2 and row[2].value else ''
        
        if name:
            # Используем factory_barcode как ключ для проверки дубликатов
            key = str(factory).strip() if factory else str(internal).strip() if internal else name.strip()
            excel_products[key] = {
                'name': str(name).strip(),
                'factory': str(factory).strip() if factory else None,
                'internal': str(internal).strip() if internal else None
            }
    
    print(f"✅ В Excel найдено товаров: {len(excel_products)}")
    
    # Загружаем из базы
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()
    
    cursor.execute('SELECT product_name, factory_barcode, internal_barcode FROM barcodes WHERE is_active = 1')
    db_products = {}
    for row in cursor.fetchall():
        # Ключ по factory_barcode или internal_barcode
        key = row['factory_barcode'] if row['factory_barcode'] else row['internal_barcode']
        if key:
            db_products[key] = {
                'name': row['product_name'],
                'factory': row['factory_barcode'],
                'internal': row['internal_barcode']
            }
    
    print(f"✅ В базе найдено штрих-кодов: {len(db_products)}")
    
    # Находим недостающие
    missing = []
    for key, product in excel_products.items():
        if key not in db_products:
            missing.append(product)
        else:
            # Проверяем, не изменилось ли название
            db_name = db_products[key]['name']
            excel_name = product['name']
            if db_name != excel_name:
                print(f"⚠️  Найдено расхождение для {key}:")
                print(f"   В базе: {db_name}")
                print(f"   В Excel: {excel_name}")
    
    print(f"\n🔍 Найдено недостающих товаров: {len(missing)}")
    
    if not missing:
        print("✅ Все товары уже загружены!")
        return
    
    # Импортируем недостающие
    print(f"\n📥 Импортируем {len(missing)} товаров...")
    
    added = 0
    skipped = 0
    
    for product in missing:
        try:
            cursor.execute('''
                INSERT INTO barcodes (product_name, factory_barcode, internal_barcode, created_by, created_at, is_active)
                VALUES (?, ?, ?, 1, datetime('now'), 1)
            ''', (product['name'], product['factory'], product['internal']))
            added += 1
            print(f"  ✅ Добавлен: {product['name'][:60]}")
        except Exception as e:
            skipped += 1
            print(f"  ❌ Пропущен: {product['name'][:60]} - {e}")
    
    conn.commit()
    conn.close()
    
    print(f"\n{'='*60}")
    print(f"✅ ИМПОРТ ЗАВЕРШЁН")
    print(f"{'='*60}")
    print(f"  Добавлено: {added}")
    print(f"  Пропущено: {skipped}")
    print(f"  Всего в базе стало: {len(db_products) + added}")

if __name__ == '__main__':
    main()
