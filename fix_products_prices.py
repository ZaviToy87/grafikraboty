# -*- coding: utf-8 -*-
"""
fix_products_prices.py — Исправление цен в базе products_1c
"""
import sqlite3
from openpyxl import load_workbook
import os
import re

DB_PATH = 'schedule.db'
EXCEL_DIR = os.path.dirname(os.path.abspath(__file__))

def normalize_name(name):
    """Нормализовать название для сравнения"""
    if not name:
        return ''
    name = str(name).strip().lower()
    name = re.sub(r'\s+', ' ', name)
    name = re.sub(r'[^\w\sа-яё0-9]', '', name)
    return name

def fix_prices():
    """Обновить цены из Excel файла"""
    print("=" * 70)
    print("🔧 ИСПРАВЛЕНИЕ ЦЕН В БАЗЕ products_1c")
    print("=" * 70)
    
    price_file = os.path.join(EXCEL_DIR, 'НОМЕНКЛАТУРА И ЦЕНА.xlsx')
    
    if not os.path.exists(price_file):
        print(f"❌ Файл не найден: {price_file}")
        return
    
    print(f"📖 Чтение файла цен: {price_file}")
    wb = load_workbook(price_file, read_only=True, data_only=True)
    ws = wb.active
    
    price_by_name = {}
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        try:
            # Пропускаем заголовок
            if row_idx == 1:
                continue
                
            # A - Наименование, F - розничная, H - Артикул
            name = str(row[0]).strip() if row[0] else None
            price = float(row[5]) if len(row) > 5 and row[5] else 0
            vendor_code = str(row[7]).strip() if len(row) > 7 and row[7] else None
            
            if name and price:
                norm_name = normalize_name(name)
                price_by_name[norm_name] = {
                    'price': price,
                    'vendor_code': vendor_code,
                    'original_name': name
                }
        except Exception as e:
            continue
    
    wb.close()
    print(f"✅ Прочитано цен: {len(price_by_name)}")
    
    # Обновляем базу данных
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    
    cursor.execute('SELECT id, name, retail_price FROM products_1c')
    products = cursor.fetchall()
    
    updated = 0
    not_found = 0
    
    print("\n🔄 Обновление цен...")
    for prod_id, name, current_price in products:
        norm_name = normalize_name(name)
        price_info = price_by_name.get(norm_name)
        
        if price_info and price_info['price'] > 0:
            # Обновляем цену
            cursor.execute('''
                UPDATE products_1c
                SET retail_price = ?, vendor_code = ?, updated_at = datetime('now')
                WHERE id = ?
            ''', (price_info['price'], price_info['vendor_code'], prod_id))
            updated += 1
            
            if updated <= 5:
                print(f"  ✅ {name[:50]}: {current_price} → {price_info['price']} ₽")
        else:
            not_found += 1
    
    conn.commit()
    conn.close()
    
    print(f"\n{'=' * 70}")
    print(f"✅ ИСПРАВЛЕНИЕ ЗАВЕРШЕНО")
    print(f"{'=' * 70}")
    print(f"📦 Всего товаров: {len(products)}")
    print(f"🔄 Обновлено цен: {updated}")
    print(f"❌ Не найдено: {not_found}")
    print(f"{'=' * 70}")

if __name__ == '__main__':
    fix_prices()
